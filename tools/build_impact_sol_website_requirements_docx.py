#!/usr/bin/env python3
"""Build the Impact Sol website vendor brief as a polished DOCX."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


INK = "111814"
MUTED = "58635D"
GREEN = "0B4A38"
GREEN_2 = "177259"
PALE = "F2F7F4"
PALE_2 = "E7F1EC"
WHITE = "FFFFFF"
BORDER = "D9E0DC"
GOLD = "C98A16"
FONT = "Arial"
MONO_FONT = "Liberation Mono"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=85, start=100, bottom=85, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="5") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_row_repeat(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def set_row_no_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_width(cell, inches: float) -> None:
    cell.width = Inches(inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_layout(table) -> None:
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def set_run_font(run, name=FONT, size=None, color=None, bold=None, italic=None) -> None:
    run.font.name = name
    if run._element.get_or_add_rPr().rFonts is None:
        run._element.get_or_add_rPr().append(OxmlElement("w:rFonts"))
    fonts = run._element.get_or_add_rPr().rFonts
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:eastAsia"), "Noto Sans Bengali")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_hyperlink(paragraph, text: str, url: str, color=GREEN_2) -> None:
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), FONT)
    r_fonts.set(qn("w:hAnsi"), FONT)
    r_pr.append(r_fonts)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    r_pr.append(c)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    run.append(r_pr)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_PATTERN = re.compile(
    r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\(https?://[^)]+\))"
)


def add_inline(paragraph, text: str, *, base_size=10.5, base_color=INK, base_bold=False) -> None:
    position = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position : match.start()])
            set_run_font(run, size=base_size, color=base_color, bold=base_bold)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=base_size, color=base_color, bold=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name=MONO_FONT, size=max(base_size - 0.7, 8), color=INK)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), PALE_2)
            run._element.get_or_add_rPr().append(shading)
        else:
            label, url = re.match(r"\[([^\]]+)\]\((https?://[^)]+)\)", token).groups()
            add_hyperlink(paragraph, label, url)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_run_font(run, size=base_size, color=base_color, bold=base_bold)


def strip_md(text: str) -> str:
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"`(.*?)`", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1", text)
    return text.strip()


def set_paragraph_base(paragraph, after=5, before=0, line=1.12) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_after = Pt(after)
    fmt.space_before = Pt(before)
    fmt.line_spacing = line
    fmt.widow_control = True


def add_body(doc, text: str, *, style=None, after=5, before=0) -> None:
    paragraph = doc.add_paragraph(style=style)
    set_paragraph_base(paragraph, after=after, before=before)
    add_inline(paragraph, text)


def add_bullet(doc, text: str, level=0) -> None:
    paragraph = doc.add_paragraph()
    fmt = paragraph.paragraph_format
    fmt.left_indent = Inches(0.25 + level * 0.22)
    fmt.first_line_indent = Inches(-0.16)
    fmt.space_after = Pt(2.5)
    fmt.line_spacing = 1.08
    fmt.widow_control = True
    marker = paragraph.add_run("•  ")
    set_run_font(marker, size=10.5, color=GREEN_2, bold=True)
    add_inline(paragraph, text)


def add_numbered(doc, marker: str, text: str) -> None:
    paragraph = doc.add_paragraph()
    fmt = paragraph.paragraph_format
    fmt.left_indent = Inches(0.34)
    fmt.first_line_indent = Inches(-0.28)
    fmt.space_after = Pt(3)
    fmt.line_spacing = 1.08
    fmt.widow_control = True
    run = paragraph.add_run(f"{marker} ")
    set_run_font(run, size=10.5, color=GREEN, bold=True)
    add_inline(paragraph, text)


def add_quote(doc, text: str) -> None:
    paragraph = doc.add_paragraph()
    fmt = paragraph.paragraph_format
    fmt.left_indent = Inches(0.35)
    fmt.right_indent = Inches(0.2)
    fmt.space_before = Pt(4)
    fmt.space_after = Pt(8)
    fmt.line_spacing = 1.12
    add_inline(paragraph, text, base_size=10.5, base_color=MUTED)
    for run in paragraph.runs:
        run.italic = True


def parse_table_row(line: str) -> list[str]:
    return [part.strip() for part in line.strip().strip("|").split("|")]


def column_widths(rows: list[list[str]], available=7.06) -> list[float]:
    count = len(rows[0])
    max_lengths = []
    for index in range(count):
        lengths = [min(max(len(strip_md(row[index])) if index < len(row) else 0, 6), 90) for row in rows]
        max_lengths.append(max(lengths) if lengths else 10)
    minimum = 0.75 if count >= 4 else 1.0
    weights = [max(length ** 0.70, 4) for length in max_lengths]
    raw = [available * weight / sum(weights) for weight in weights]
    raw = [max(minimum, width) for width in raw]
    scale = available / sum(raw)
    widths = [round(width * scale, 2) for width in raw]
    widths[-1] += round(available - sum(widths), 2)
    return widths


def add_table(doc, rows: list[list[str]]) -> None:
    if not rows:
        return
    columns = len(rows[0])
    normalized = [row + [""] * (columns - len(row)) for row in rows]
    widths = column_widths(normalized)
    table = doc.add_table(rows=len(normalized), cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_repeat_table_layout(table)
    set_table_borders(table)
    for r_index, source_row in enumerate(normalized):
        row = table.rows[r_index]
        set_row_no_split(row)
        if r_index == 0:
            set_row_repeat(row)
        for c_index, value in enumerate(source_row):
            cell = row.cells[c_index]
            set_cell_width(cell, widths[c_index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if r_index == 0:
                set_cell_shading(cell, GREEN)
            elif r_index % 2 == 0:
                set_cell_shading(cell, PALE)
            else:
                set_cell_shading(cell, WHITE)
            paragraph = cell.paragraphs[0]
            set_paragraph_base(paragraph, after=0, line=1.04)
            add_inline(
                paragraph,
                value,
                base_size=8.4 if columns >= 4 else 8.8,
                base_color=WHITE if r_index == 0 else INK,
                base_bold=r_index == 0,
            )
    following = doc.add_paragraph()
    following.paragraph_format.space_after = Pt(1)
    following.paragraph_format.line_spacing = 0.2


def style_document(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans Bengali")
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.widow_control = True

    title = styles["Title"]
    title.font.name = FONT
    title.font.size = Pt(31)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string("000000")
    title._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    title._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    title_p_pr = title._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    for style_name, size, before, after in (
        ("Heading 1", 19, 18, 7),
        ("Heading 2", 14, 13, 5),
        ("Heading 3", 11.5, 9, 3),
    ):
        style = styles[style_name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string("000000")
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    if "Small Label" not in styles:
        label = styles.add_style("Small Label", WD_STYLE_TYPE.PARAGRAPH)
    else:
        label = styles["Small Label"]
    label.font.name = FONT
    label.font.size = Pt(9)
    label.font.bold = True
    label.font.color.rgb = RGBColor.from_string(GREEN)
    label.paragraph_format.space_after = Pt(17)
    label._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    label._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)


def set_section(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.67)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)
    section.different_first_page_header_footer = True


def add_page_field(paragraph) -> None:
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.append(fld_char_1)
    run._r.append(instr_text)
    run._r.append(fld_char_2)
    set_run_font(run, size=8.5, color=MUTED)


def add_footer(section) -> None:
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        Inches(7.0), WD_TAB_ALIGNMENT.RIGHT
    )
    run = paragraph.add_run("Impact Sol Website Requirements")
    set_run_font(run, size=8.2, color=MUTED)
    run = paragraph.add_run("\t")
    set_run_font(run, size=8.2, color=MUTED)
    add_page_field(paragraph)
    first = section.first_page_footer
    first.paragraphs[0].text = ""


def extract_meta(lines: list[str]) -> tuple[str, list[tuple[str, str]], int]:
    title = strip_md(lines[0].lstrip("# "))
    metadata: list[tuple[str, str]] = []
    index = 1
    for index in range(1, len(lines)):
        line = lines[index].strip()
        if line.startswith("## "):
            break
        if not line:
            continue
        match = re.match(r"\*\*([^*]+):\*\*\s*(.*?)(?:\s{2})?$", line)
        if match:
            metadata.append((match.group(1), strip_md(match.group(2))))
    return title, metadata, index


def add_cover(doc: Document, title_text: str, metadata: list[tuple[str, str]]) -> None:
    label = doc.add_paragraph(style="Small Label")
    label.add_run("VENDOR REQUIREMENTS BRIEF")

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(16)
    title.add_run(title_text)

    purpose = next((value for key, value in metadata if key == "Purpose"), "")
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(34)
    subtitle.paragraph_format.line_spacing = 1.18
    run = subtitle.add_run(purpose)
    set_run_font(run, size=14, color=MUTED)

    deck = doc.add_paragraph()
    deck.paragraph_format.space_after = Pt(34)
    deck.paragraph_format.line_spacing = 1.2
    run = deck.add_run(
        "A complete handoff for strategy, content, user experience, design, "
        "development, quality assurance, launch, and long-term ownership."
    )
    set_run_font(run, size=11.5, color=INK)

    for key, value in metadata:
        if key == "Purpose":
            continue
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(5)
        key_run = paragraph.add_run(f"{key}: ")
        set_run_font(key_run, size=9.5, color=GREEN, bold=True)
        value_run = paragraph.add_run(value)
        set_run_font(value_run, size=9.5, color=INK)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(22)
    run = spacer.add_run("IMPACT SOL.")
    set_run_font(run, size=10, color=GREEN_2, bold=True)
    doc.add_page_break()


def add_contents(doc: Document, lines: list[str]) -> None:
    heading = doc.add_paragraph(style="Heading 1")
    heading.add_run("Contents")
    heading.paragraph_format.space_before = Pt(0)
    intro = doc.add_paragraph()
    set_paragraph_base(intro, after=12)
    add_inline(
        intro,
        "Use this brief as the scope baseline. The vendor should answer it with an item-by-item compliance matrix and identify every assumption, alternative, exclusion, and recurring cost.",
    )
    for line in lines:
        if not line.startswith("## "):
            continue
        value = strip_md(line[3:])
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.08)
        paragraph.paragraph_format.space_after = Pt(3.2)
        paragraph.paragraph_format.keep_together = True
        run = paragraph.add_run(value)
        set_run_font(run, size=9.8, color=INK, bold=value == "Executive brief")
    doc.add_page_break()


def parse_markdown(doc: Document, lines: list[str], start: int) -> None:
    index = start
    while index < len(lines):
        raw = lines[index].rstrip()
        line = raw.strip()
        if not line:
            index += 1
            continue
        if line.startswith("## "):
            paragraph = doc.add_paragraph(style="Heading 1")
            paragraph.add_run(strip_md(line[3:]))
            index += 1
            continue
        if line.startswith("### "):
            paragraph = doc.add_paragraph(style="Heading 2")
            paragraph.add_run(strip_md(line[4:]))
            index += 1
            continue
        if line.startswith("#### "):
            paragraph = doc.add_paragraph(style="Heading 3")
            paragraph.add_run(strip_md(line[5:]))
            index += 1
            continue
        if line.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                candidate = parse_table_row(lines[index])
                if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in candidate):
                    rows.append(candidate)
                index += 1
            add_table(doc, rows)
            continue
        if line.startswith("> "):
            parts = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                parts.append(lines[index].strip().lstrip("> "))
                index += 1
            add_quote(doc, " ".join(parts))
            continue
        bullet = re.match(r"^-\s+(.*)$", line)
        if bullet:
            add_bullet(doc, bullet.group(1))
            index += 1
            continue
        numbered = re.match(r"^(\d+\.)\s+(.*)$", line)
        if numbered:
            add_numbered(doc, numbered.group(1), numbered.group(2))
            index += 1
            continue
        paragraph_parts = [line]
        index += 1
        while index < len(lines):
            next_line = lines[index].strip()
            if not next_line:
                break
            if (
                next_line.startswith(("## ", "### ", "#### ", "|", "> ", "- "))
                or re.match(r"^\d+\.\s+", next_line)
            ):
                break
            paragraph_parts.append(next_line)
            index += 1
        add_body(doc, " ".join(paragraph_parts))


def build(source: Path, output: Path) -> None:
    text = source.read_text(encoding="utf-8")
    lines = text.splitlines()
    title_text, metadata, start = extract_meta(lines)

    doc = Document()
    style_document(doc)
    section = doc.sections[0]
    set_section(section)
    add_footer(section)

    doc.core_properties.title = title_text
    doc.core_properties.subject = "Website requirements and vendor handoff brief"
    doc.core_properties.author = "Impact Sol."
    doc.core_properties.keywords = "Impact Sol, Chokro, website requirements, vendor brief"
    doc.core_properties.comments = "Prepared for vendor scoping and implementation."

    add_cover(doc, title_text, metadata)
    add_contents(doc, lines[start:])
    parse_markdown(doc, lines, start)

    for paragraph in doc.paragraphs:
        paragraph.paragraph_format.widow_control = True
        if paragraph.style.name.startswith("Heading"):
            paragraph.paragraph_format.keep_with_next = True

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> int:
    if len(sys.argv) != 3:
        print("Usage: build_impact_sol_website_requirements_docx.py SOURCE.md OUTPUT.docx")
        return 2
    build(Path(sys.argv[1]).resolve(), Path(sys.argv[2]).resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
